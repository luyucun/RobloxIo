from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "IO_BaseBalanceDraft.xlsx"
CONFIG_PATH = ROOT / "CodeConfig.lua"
SHEET_NAME = "兑换码"
HEADER_ROW = 6
DATA_START_ROW = 7
BEGIN_MARKER = "-- BEGIN GENERATED CODE ROWS"
END_MARKER = "-- END GENERATED CODE ROWS"

ONLINE_REWARD_CONFIG_PATH = ROOT / "OnlineRewardConfig.lua"
ONLINE_REWARD_SHEET_NAME = "在线奖励"
ONLINE_REWARD_HEADER_ROW = 9
ONLINE_REWARD_DATA_START_ROW = 10
ONLINE_REWARD_BEGIN_MARKER = "-- BEGIN GENERATED ONLINE REWARD ROWS"
ONLINE_REWARD_END_MARKER = "-- END GENERATED ONLINE REWARD ROWS"

SEVEN_DAY_LOGIN_REWARD_CONFIG_PATH = ROOT / "SevenDayLoginRewardConfig.lua"
SEVEN_DAY_LOGIN_REWARD_SHEET_NAME = "七日登录奖励"
SEVEN_DAY_LOGIN_REWARD_BEGIN_MARKER = "-- BEGIN GENERATED SEVEN DAY LOGIN REWARD ROWS"
SEVEN_DAY_LOGIN_REWARD_END_MARKER = "-- END GENERATED SEVEN DAY LOGIN REWARD ROWS"

SKIN_CONFIG_PATH = ROOT / "SkinConfig.lua"
SKIN_BEGIN_MARKER = "-- BEGIN GENERATED SKIN ROWS"
SKIN_END_MARKER = "-- END GENERATED SKIN ROWS"

TRAIL_CONFIG_PATH = ROOT / "TrailConfig.lua"
TRAIL_SHEET_NAME = "尾迹"
TRAIL_HEADER_ROW = 6
TRAIL_DATA_START_ROW = 7
TRAIL_BEGIN_MARKER = "-- BEGIN GENERATED TRAIL ROWS"
TRAIL_END_MARKER = "-- END GENERATED TRAIL ROWS"

TITLE_CONFIG_PATH = ROOT / "TitleConfig.lua"
TITLE_SHEET_NAME = "称号"
TITLE_HEADER_ROW = 4
TITLE_DATA_START_ROW = 5
TITLE_BEGIN_MARKER = "-- BEGIN GENERATED TITLE ROWS"
TITLE_END_MARKER = "-- END GENERATED TITLE ROWS"

SHOP_CONFIG_PATH = ROOT / "ShopConfig.lua"
DIAMOND_SHOP_SHEET_NAME = "钻石购买"
DIAMOND_SHOP_HEADER_ROW = 5
DIAMOND_SHOP_DATA_START_ROW = 6
DIAMOND_SHOP_BEGIN_MARKER = "-- BEGIN GENERATED DIAMOND PRODUCT ROWS"
DIAMOND_SHOP_END_MARKER = "-- END GENERATED DIAMOND PRODUCT ROWS"

POTION_CONFIG_PATH = ROOT / "PotionConfig.lua"
POTION_SHEET_NAME = "药水"
POTION_HEADER_ROW = 2
POTION_DATA_START_ROW = 3
POTION_BEGIN_MARKER = "-- BEGIN GENERATED POTION ROWS"
POTION_END_MARKER = "-- END GENERATED POTION ROWS"

WHEEL_CONFIG_PATH = ROOT / "WheelConfig.lua"
WHEEL_SHEET_NAME = "转盘规划"
WHEEL_HEADER_ROW = 6
WHEEL_DATA_START_ROW = 7
WHEEL_BEGIN_MARKER = "-- BEGIN GENERATED WHEEL REWARD ROWS"
WHEEL_END_MARKER = "-- END GENERATED WHEEL REWARD ROWS"

ATTRIBUTE_CONFIG_PATH = ROOT / "AttributeConfig.lua"
ATTRIBUTE_CONFIG_SHEET_NAME = "属性养成配置"
ATTRIBUTE_PRODUCT_SHEET_NAME = "属性养成新的开发者商品"
ATTRIBUTE_CONFIG_HEADER_ROW = 4
ATTRIBUTE_CONFIG_DATA_START_ROW = 5
ATTRIBUTE_PRODUCT_HEADER_ROW = 12
ATTRIBUTE_PRODUCT_DATA_START_ROW = 13
ATTRIBUTE_CONFIG_BEGIN_MARKER = "-- BEGIN GENERATED ATTRIBUTE CONFIG ROWS"
ATTRIBUTE_CONFIG_END_MARKER = "-- END GENERATED ATTRIBUTE CONFIG ROWS"

TASK_CONFIG_PATH = ROOT / "TaskConfig.lua"
TASK_SHEET_NAME = "任务系统数据表"
TASK_HEADER_ROW = 21
TASK_DATA_START_ROW = 22
TASK_BEGIN_MARKER = "-- BEGIN GENERATED TASK ROWS"
TASK_END_MARKER = "-- END GENERATED TASK ROWS"


def is_blank(value) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"null", "none", "-"}


def lua_string(value) -> str:
    text = str(value)
    if "'" not in text and "\\" not in text and "\n" not in text:
        return "'" + text + "'"
    for equals_count in range(12):
        equals = "=" * equals_count
        close = f"]{equals}]"
        if close not in text:
            return f"[{equals}[{text}]{equals}]"
    raise RuntimeError("Could not encode Lua string safely.")


def lua_value(value) -> str:
    if is_blank(value):
        return "nil"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(int(value)) if float(value).is_integer() else repr(float(value))
    return lua_string(value)


def parse_potion_ids() -> list[int]:
    source = (ROOT / "PotionConfig.lua").read_text(encoding="utf-8")
    match = re.search(r"PotionConfig\.OrderedPotionIds\s*=\s*{(?P<body>.*?)\n}", source, re.S)
    if not match:
        return []
    return [int(value) for value in re.findall(r"\b\d+\b", match.group("body"))]


POTION_IDS = parse_potion_ids()
POTION_NAME_ALIASES = {
    "加初级药水": 1001,
    "初级药水": 1001,
    "加基础药水": 1001,
    "基础药水": 1001,
    "BasicPotion": 1001,
    "加高级药水": 1002,
    "高级药水": 1002,
    "AdvancedPotion": 1002,
    "加稀有药水": 1003,
    "稀有药水": 1003,
    "RarePotion": 1003,
}


def parse_amount_token(token: str | None) -> int:
    if is_blank(token):
        return 1
    return max(1, int(float(str(token).strip())))


def parse_potion_reward(compact: str) -> dict | None:
    body = None
    for prefix in ("药水", "Potion", "potion"):
        if compact.startswith(prefix):
            body = compact[len(prefix) :]
            break
    if not body:
        return None

    explicit = re.fullmatch(r"(\d+)(?:[xX*×+](\d+))?", body)
    if explicit and explicit.group(2):
        return {
            "RewardType": "Potion",
            "PotionId": int(explicit.group(1)),
            "Amount": parse_amount_token(explicit.group(2)),
        }

    best_id = None
    best_suffix = None
    for potion_id in sorted(POTION_IDS, key=lambda value: len(str(value)), reverse=True):
        potion_id_text = str(potion_id)
        if body == potion_id_text:
            best_id = potion_id
            best_suffix = None
            break
        if body.startswith(potion_id_text):
            suffix = body[len(potion_id_text) :]
            if suffix.isdigit():
                best_id = potion_id
                best_suffix = suffix
                break

    if best_id is None and explicit:
        best_id = int(explicit.group(1))

    if best_id is None:
        return None

    return {
        "RewardType": "Potion",
        "PotionId": best_id,
        "Amount": parse_amount_token(best_suffix),
    }


def parse_reward_type_and_amount(reward_type_value, amount_value) -> tuple[dict | None, str | None]:
    if is_blank(reward_type_value):
        return None, None

    reward_type_text = str(reward_type_value).strip()
    compact = re.sub(r"\s+", "", reward_type_text).replace("＋", "+")
    amount = parse_amount_token(amount_value)

    if compact in {"经验值", "加经验", "加经验值", "Experience", "EXP", "Exp"}:
        return {"RewardType": "Experience", "Amount": amount}, None

    if compact in {"加钻石", "钻石", "Diamonds", "Diamond"}:
        return {"RewardType": "Diamonds", "Amount": amount}, None

    if compact in {"加转盘次数", "转盘次数", "转盘", "WheelSpins", "Spins"}:
        return {"RewardType": "WheelSpins", "Amount": amount}, None

    if compact in {"加护盾", "护盾", "Shield"}:
        return {"RewardType": "Shield", "Amount": amount, "DurationSeconds": amount}, None

    if compact in {"特殊皮肤", "皮肤", "Skin", "SpecialSkin"}:
        return {"RewardType": "Skin", "SkinId": amount, "Amount": 1}, None

    potion = parse_potion_reward(compact)
    if potion:
        potion["Amount"] = amount
        return potion, None

    alias_potion_id = POTION_NAME_ALIASES.get(compact)
    if alias_potion_id:
        return {"RewardType": "Potion", "PotionId": alias_potion_id, "Amount": amount}, None

    match = re.search(r"(\d+)", compact)
    if "药水" in compact or "Potion" in compact:
        if match:
            return {"RewardType": "Potion", "PotionId": int(match.group(1)), "Amount": amount}, None
        return None, f"无法解析在线奖励药水ID: {reward_type_text}"

    return None, f"无法解析在线奖励类型: {reward_type_text}"


def parse_reward_text(value) -> tuple[dict | None, str | None]:
    if is_blank(value):
        return None, None

    text = str(value).strip()
    compact = re.sub(r"\s+", "", text)
    compact = compact.replace("＋", "+")

    diamond = re.fullmatch(r"(?:加)?钻石[+]?(\d+)", compact) or re.fullmatch(r"Diamonds[+]?(\d+)", compact)
    if diamond:
        return {"RewardType": "Diamonds", "Amount": parse_amount_token(diamond.group(1))}, None

    wheel = (
        re.fullmatch(r"(?:加)?转盘次数[+]?(\d+)", compact)
        or re.fullmatch(r"转盘[+]?(\d+)", compact)
        or re.fullmatch(r"WheelSpins[+]?(\d+)", compact)
        or re.fullmatch(r"Spins[+]?(\d+)", compact)
    )
    if wheel:
        return {"RewardType": "WheelSpins", "Amount": parse_amount_token(wheel.group(1))}, None

    shield = re.fullmatch(r"(\d+)秒护盾", compact) or re.fullmatch(r"Shield(\d+)s?", compact)
    if shield:
        duration_seconds = parse_amount_token(shield.group(1))
        return {"RewardType": "Shield", "Amount": duration_seconds, "DurationSeconds": duration_seconds}, None

    skin = re.fullmatch(r"(?:特殊)?(?:武器)?皮肤(\d+)", compact) or re.fullmatch(r"(?:Special)?(?:Weapon)?Skin(\d+)", compact)
    if skin:
        return {"RewardType": "PendingWeaponSkin", "SkinId": parse_amount_token(skin.group(1)), "Amount": 1, "Pending": True}, None

    potion = parse_potion_reward(compact)
    if potion:
        return potion, None

    alias_potion_id = POTION_NAME_ALIASES.get(compact)
    if alias_potion_id:
        return {"RewardType": "Potion", "PotionId": alias_potion_id, "Amount": 1}, None

    for alias, potion_id in sorted(POTION_NAME_ALIASES.items(), key=lambda item: len(item[0]), reverse=True):
        if compact.startswith(alias):
            suffix = compact[len(alias):]
            amount_match = re.fullmatch(r"[+xX*×]?(\d+)", suffix)
            if amount_match:
                return {"RewardType": "Potion", "PotionId": potion_id, "Amount": parse_amount_token(amount_match.group(1))}, None

    return None, f"无法解析奖励内容: {text}"


def format_reward(reward: dict) -> str:
    reward_type = reward["RewardType"]
    if reward_type == "Potion":
        return "{ RewardType = 'Potion', PotionId = %d, Amount = %d }" % (
            int(reward["PotionId"]),
            int(reward["Amount"]),
        )
    if reward_type == "WheelSpins":
        return "{ RewardType = 'WheelSpins', Amount = %d }" % int(reward["Amount"])
    if reward_type == "Diamonds":
        return "{ RewardType = 'Diamonds', Amount = %d }" % int(reward["Amount"])
    raise RuntimeError(f"Unsupported reward type: {reward_type}")


def read_code_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = workbook[SHEET_NAME]
    rows = []
    warnings = []

    for row_index in range(DATA_START_ROW, worksheet.max_row + 1):
        code_id = worksheet.cell(row_index, 3).value
        code_text = worksheet.cell(row_index, 4).value
        code_type = worksheet.cell(row_index, 5).value
        expire_at = worksheet.cell(row_index, 6).value
        max_uses = worksheet.cell(row_index, 7).value

        if all(is_blank(value) for value in (code_id, code_text, code_type, expire_at, max_uses)):
            reward_cells = [worksheet.cell(row_index, column).value for column in range(8, 11)]
            if all(is_blank(value) for value in reward_cells):
                continue

        rewards = []
        for column in range(8, 11):
            reward, warning = parse_reward_text(worksheet.cell(row_index, column).value)
            if reward:
                rewards.append(reward)
            if warning:
                warnings.append({"row": row_index, "column": column, "warning": warning})

        rows.append(
            {
                "Row": row_index,
                "CodeId": 0 if is_blank(code_id) else int(float(code_id)),
                "CodeText": "" if is_blank(code_text) else str(code_text).strip(),
                "CodeType": "Timed" if is_blank(code_type) else str(code_type).strip(),
                "ExpireAt": None if is_blank(expire_at) else expire_at,
                "MaxUses": None if is_blank(max_uses) else int(float(max_uses)),
                "Rewards": rewards,
            }
        )

    return rows, warnings


def build_generated_block(rows) -> str:
    lines = [
        BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 兑换码. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "CodeConfig.ExcelRows = {",
    ]
    for row in rows:
        lines.extend(
            [
                "    {",
                f"        Row = {row['Row']},",
                f"        ['兑换码ID'] = {row['CodeId']},",
                f"        ['兑换码文本'] = {lua_value(row['CodeText'])},",
                f"        ['类型'] = {lua_value(row['CodeType'])},",
                f"        ['失效时间'] = {lua_value(row['ExpireAt'])},",
                f"        ['使用人数上限'] = {lua_value(row['MaxUses'])},",
                "        Rewards = {",
            ]
        )
        for reward in row["Rewards"]:
            lines.append(f"            {format_reward(reward)},")
        lines.extend(["        },", "    },"])
    lines.extend(["}", END_MARKER])
    return "\n".join(lines)


def build_online_reward_row(reward: dict) -> str:
    fields = [
        f"Id = {int(reward['Id'])}",
        f"RewardType = {lua_value(reward['RewardType'])}",
    ]
    if reward.get("PotionId"):
        fields.append(f"PotionId = {int(reward['PotionId'])}")
    if reward.get("SkinId"):
        fields.append(f"SkinId = {int(reward['SkinId'])}")
    if not is_blank(reward.get("Label")):
        fields.append(f"Label = {lua_value(reward['Label'])}")
    fields.append(f"Amount = {int(reward['Amount'])}")
    if reward.get("DurationSeconds"):
        fields.append(f"DurationSeconds = {int(reward['DurationSeconds'])}")
    fields.extend(
        [
            f"RequiredSeconds = {int(reward['RequiredSeconds'])}",
            f"Icon = {lua_value(reward['Icon'])}",
        ]
    )
    return "    { " + ", ".join(fields) + " },"


def build_online_reward_generated_block(rows) -> str:
    lines = [
        ONLINE_REWARD_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 在线奖励. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "OnlineRewardConfig.ExcelRows = {",
    ]
    for reward in rows:
        lines.append(build_online_reward_row(reward))
    lines.extend(["}", ONLINE_REWARD_END_MARKER])
    return "\n".join(lines)


def read_online_reward_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = workbook[ONLINE_REWARD_SHEET_NAME]
    rows = []
    warnings = []

    for row_index in range(ONLINE_REWARD_DATA_START_ROW, worksheet.max_row + 1):
        reward_id = worksheet.cell(row_index, 5).value
        reward_type = worksheet.cell(row_index, 6).value
        reward_label = worksheet.cell(row_index, 7).value
        reward_amount = worksheet.cell(row_index, 8).value
        reward_icon = worksheet.cell(row_index, 9).value
        required_seconds = worksheet.cell(row_index, 10).value

        if all(is_blank(value) for value in (reward_id, reward_type, reward_label, reward_amount, reward_icon, required_seconds)):
            continue

        reward, warning = parse_reward_type_and_amount(reward_type, reward_amount)
        if warning:
            warnings.append({"row": row_index, "column": 6, "warning": warning})
        if not reward:
            continue

        reward["Id"] = int(float(reward_id)) if not is_blank(reward_id) else len(rows) + 1
        reward["Label"] = "" if is_blank(reward_label) else str(reward_label).strip()
        reward["RequiredSeconds"] = int(float(required_seconds)) if not is_blank(required_seconds) else 0
        reward["Icon"] = "" if is_blank(reward_icon) else str(reward_icon).strip()
        rows.append(reward)

    return rows, warnings


def get_sheet(workbook, sheet_name: str, fallback_index: int):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.worksheets[fallback_index]


def read_weapon_skin_metadata(workbook) -> dict[int, dict]:
    worksheet = get_sheet(workbook, "武器数值", 4)
    metadata = {}
    for row_index in range(5, worksheet.max_row + 1):
        skin_id = worksheet.cell(row_index, 1).value
        is_skin = worksheet.cell(row_index, 3).value
        if math_safe_int(is_skin, 0) != 1:
            continue
        resolved_skin_id = math_safe_int(skin_id, 0)
        if resolved_skin_id <= 0:
            continue
        template_name = worksheet.cell(row_index, 4).value
        skin_name = worksheet.cell(row_index, 5).value
        icon = worksheet.cell(row_index, 6).value
        template_path = worksheet.cell(row_index, 7).value
        metadata[resolved_skin_id] = {
            "Name": "" if is_blank(skin_name) else str(skin_name).strip(),
            "TemplateName": "" if is_blank(template_name) else str(template_name).strip(),
            "TemplatePath": "" if is_blank(template_path) else str(template_path).strip(),
            "IconImage": "" if is_blank(icon) else str(icon).strip(),
        }
    return metadata


def math_safe_int(value, default=0) -> int:
    if is_blank(value):
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def math_safe_float(value, default=0.0) -> float:
    if is_blank(value):
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def bool_from_cell(value, default=False) -> bool:
    if is_blank(value):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value) != 0
    text = str(value).strip().lower()
    return text in {"true", "yes", "y", "1", "启用", "是"}


def build_header_map(worksheet, header_row: int) -> dict[str, int]:
    headers = {}
    for column_index in range(1, worksheet.max_column + 1):
        value = worksheet.cell(header_row, column_index).value
        if not is_blank(value):
            headers[str(value).strip()] = column_index
    return headers


def cell_by_header(worksheet, row_index: int, headers: dict[str, int], header: str):
    column_index = headers.get(header)
    return worksheet.cell(row_index, column_index).value if column_index else None


def read_potion_rows() -> list[dict]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, POTION_SHEET_NAME, 8)
    headers = build_header_map(worksheet, POTION_HEADER_ROW)
    required_headers = [
        "ID",
        "药水名字",
        "稀有度",
        "药水图标",
        "药水模型名字",
        "药水生效时间（秒）",
        "药水经验加成比例",
        "钻石价格",
        "罗布币价格",
        "开发者商品ID",
    ]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise RuntimeError("Missing potion headers: " + ", ".join(missing_headers))

    rows = []
    for row_index in range(POTION_DATA_START_ROW, worksheet.max_row + 1):
        potion_id = math_safe_int(cell_by_header(worksheet, row_index, headers, "ID"), 0)
        if potion_id <= 0:
            continue
        rows.append({
            "Id": potion_id,
            "Name": "" if is_blank(cell_by_header(worksheet, row_index, headers, "药水名字")) else str(cell_by_header(worksheet, row_index, headers, "药水名字")).strip(),
            "Rarity": math_safe_int(cell_by_header(worksheet, row_index, headers, "稀有度"), 0),
            "IconImage": "" if is_blank(cell_by_header(worksheet, row_index, headers, "药水图标")) else str(cell_by_header(worksheet, row_index, headers, "药水图标")).strip(),
            "ModelName": "" if is_blank(cell_by_header(worksheet, row_index, headers, "药水模型名字")) else str(cell_by_header(worksheet, row_index, headers, "药水模型名字")).strip(),
            "DurationSeconds": math_safe_int(cell_by_header(worksheet, row_index, headers, "药水生效时间（秒）"), 0),
            "ExperienceBonus": math_safe_float(cell_by_header(worksheet, row_index, headers, "药水经验加成比例"), 0),
            "MoveSpeedBonus": math_safe_float(cell_by_header(worksheet, row_index, headers, "药水移动速度加成比例"), 0),
            "DiamondPrice": math_safe_int(cell_by_header(worksheet, row_index, headers, "钻石价格"), 0),
            "RobuxPrice": math_safe_int(cell_by_header(worksheet, row_index, headers, "罗布币价格"), 0),
            "ProductId": math_safe_int(cell_by_header(worksheet, row_index, headers, "开发者商品ID"), 0),
        })
    rows.sort(key=lambda row: row["Id"])
    return rows


def build_potion_generated_block(rows: list[dict]) -> str:
    lines = [
        POTION_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 药水. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "PotionConfig.OrderedPotionIds = {",
    ]
    for row in rows:
        lines.append(f"    {row['Id']},")
    lines.extend(["}", "", "PotionConfig.Potions = {"])
    for row in rows:
        lines.extend([
            f"    [{row['Id']}] = {{",
            f"        Id = {row['Id']},",
            f"        Name = {lua_value(row['Name'])},",
            f"        Rarity = {row['Rarity']},",
            f"        IconImage = {lua_value(row['IconImage'])},",
            f"        ModelName = {lua_value(row['ModelName'])},",
            f"        DurationSeconds = {row['DurationSeconds']},",
            f"        ExperienceBonus = {lua_number(row['ExperienceBonus'])},",
            f"        MoveSpeedBonus = {lua_number(row['MoveSpeedBonus'])},",
            f"        DiamondPrice = {row['DiamondPrice']},",
            f"        RobuxPrice = {row['RobuxPrice']},",
            f"        ProductId = {row['ProductId']},",
            "    },",
        ])
    lines.extend(["}", POTION_END_MARKER])
    return "\n".join(lines)


def build_wheel_reward_label(reward: dict) -> str:
    reward_type = reward["RewardType"]
    if reward_type == "Shield":
        return f"{int(reward['DurationSeconds'])}s Shield"
    if reward_type == "Potion":
        potion_names = {
            1001: "Basic Potion",
            1002: "Advanced Potion",
            1003: "Rare Potion",
        }
        return f"{potion_names.get(math_safe_int(reward.get('PotionId'), 0), 'Potion')} +{math_safe_int(reward.get('Amount'), 1)}"
    if reward_type == "PendingWeaponSkin":
        return f"Special Weapon Skin {math_safe_int(reward.get('SkinId'), 0)}"
    if reward_type == "WheelSpins":
        return f"Wheel Spins +{math_safe_int(reward.get('Amount'), 0)}"
    if reward_type == "Diamonds":
        return f"Diamonds +{math_safe_int(reward.get('Amount'), 0)}"
    return reward_type


def build_wheel_reward_id(reward: dict) -> str:
    reward_type = reward["RewardType"]
    if reward_type == "Shield":
        return f"Shield{math_safe_int(reward.get('DurationSeconds'), 0)}"
    if reward_type == "Potion":
        return f"Potion{math_safe_int(reward.get('PotionId'), 0)}"
    if reward_type == "PendingWeaponSkin":
        return f"WeaponSkin{math_safe_int(reward.get('SkinId'), 0)}"
    if reward_type == "WheelSpins":
        return f"WheelSpins{math_safe_int(reward.get('Amount'), 0)}"
    if reward_type == "Diamonds":
        return f"Diamonds{math_safe_int(reward.get('Amount'), 0)}"
    return reward_type


def read_wheel_reward_rows() -> tuple[list[dict], list[dict]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, WHEEL_SHEET_NAME, 11)
    headers = build_header_map(worksheet, WHEEL_HEADER_ROW)
    required_headers = ["转盘位置", "奖励内容", "对应权重", "对应转盘子节点", "对应旋转角度"]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise RuntimeError("Missing wheel headers: " + ", ".join(missing_headers))

    rows = []
    warnings = []
    for row_index in range(WHEEL_DATA_START_ROW, worksheet.max_row + 1):
        slot = math_safe_int(cell_by_header(worksheet, row_index, headers, "转盘位置"), 0)
        reward_text = cell_by_header(worksheet, row_index, headers, "奖励内容")
        if slot <= 0 and is_blank(reward_text):
            continue
        reward, warning = parse_reward_text(reward_text)
        if warning:
            warnings.append({"row": row_index, "column": "奖励内容", "warning": warning})
            continue
        if not reward:
            continue
        row = dict(reward)
        row.update({
            "Slot": slot,
            "Id": build_wheel_reward_id(reward),
            "Label": build_wheel_reward_label(reward),
            "Weight": math_safe_float(cell_by_header(worksheet, row_index, headers, "对应权重"), 0),
            "GiftName": "" if is_blank(cell_by_header(worksheet, row_index, headers, "对应转盘子节点")) else str(cell_by_header(worksheet, row_index, headers, "对应转盘子节点")).strip(),
            "TargetRotation": math_safe_int(cell_by_header(worksheet, row_index, headers, "对应旋转角度"), 0),
        })
        rows.append(row)
    rows.sort(key=lambda row: row["Slot"])
    return rows, warnings


def build_wheel_reward_generated_block(rows: list[dict]) -> str:
    lines = [
        WHEEL_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 转盘规划. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "WheelConfig.Rewards = {",
    ]
    for row in rows:
        lines.extend([
            "    {",
            f"        Slot = {row['Slot']},",
            f"        Id = {lua_value(row['Id'])},",
            f"        RewardType = {lua_value(row['RewardType'])},",
            f"        Label = {lua_value(row['Label'])},",
            f"        Weight = {lua_number(row['Weight'])},",
            f"        GiftName = {lua_value(row['GiftName'])},",
            f"        TargetRotation = {row['TargetRotation']},",
        ])
        if row["RewardType"] == "Shield":
            lines.append(f"        DurationSeconds = {math_safe_int(row.get('DurationSeconds'), 0)},")
        elif row["RewardType"] == "Potion":
            lines.append(f"        PotionId = {math_safe_int(row.get('PotionId'), 0)},")
            lines.append(f"        Amount = {math_safe_int(row.get('Amount'), 1)},")
        elif row["RewardType"] == "PendingWeaponSkin":
            lines.append(f"        SkinId = {math_safe_int(row.get('SkinId'), 0)},")
            lines.append("        Pending = true,")
        elif row["RewardType"] in {"WheelSpins", "Diamonds"}:
            lines.append(f"        Amount = {math_safe_int(row.get('Amount'), 0)},")
        lines.append("    },")
    lines.extend(["}", WHEEL_END_MARKER])
    return "\n".join(lines)


TASK_TYPE_BY_ID = {
    1001: ("Daily", "OnlineSeconds"),
    1002: ("Daily", "PlayerKills"),
    1003: ("Daily", "InviteFriend"),
    1004: ("Daily", "WheelSpinsUsed"),
    2001: ("Weekly", "OnlineSeconds"),
    2002: ("Weekly", "PlayerKills"),
    2003: ("Weekly", "DiamondsEarned"),
    2004: ("Weekly", "LoginDays"),
}

TASK_POTION_REWARD_IDS = {
    "中级药水": 1002,
    "高级药水": 1003,
    "AdvancedPotion": 1002,
    "RarePotion": 1003,
}


def parse_task_target(task_type: str, description: str) -> int:
    text = str(description or "")
    if task_type == "OnlineSeconds":
        hour_match = re.search(r"(\d+)\s*小时", text)
        if hour_match:
            return math_safe_int(hour_match.group(1), 1) * 3600
        minute_match = re.search(r"(\d+)\s*分钟", text)
        if minute_match:
            return math_safe_int(minute_match.group(1), 1) * 60
        english_hour = re.search(r"(\d+)\s*h(?:our)?", text, re.I)
        if english_hour:
            return math_safe_int(english_hour.group(1), 1) * 3600
        english_minute = re.search(r"(\d+)\s*m(?:in)?", text, re.I)
        if english_minute:
            return math_safe_int(english_minute.group(1), 1) * 60

    if task_type == "InviteFriend":
        match = re.search(r"(\d+)", text)
        return math_safe_int(match.group(1), 1) if match else 1

    match = re.search(r"(\d+)", text)
    if match:
        return math_safe_int(match.group(1), 1)
    return 1


def parse_task_reward(reward_value, amount_value) -> dict:
    reward_text = "" if is_blank(reward_value) else str(reward_value).strip()
    compact = re.sub(r"\s+", "", reward_text)
    amount = parse_amount_token(amount_value)

    if compact in {"钻石", "宝石", "Diamonds", "Diamond", "Gems", "Gem"}:
        return {"RewardType": "Diamonds", "Amount": amount}
    if compact in {"转盘次数", "转盘", "WheelSpins", "Spins"}:
        return {"RewardType": "WheelSpins", "Amount": amount}
    if compact in {"经验", "经验值", "Experience", "EXP", "Exp"}:
        return {"RewardType": "Experience", "Amount": amount}
    potion_id = TASK_POTION_REWARD_IDS.get(compact)
    if potion_id:
        return {"RewardType": "Potion", "PotionId": potion_id, "Amount": amount}

    reward, warning = parse_reward_type_and_amount(reward_text, amount)
    if warning or not reward:
        raise RuntimeError(f"Unable to parse task reward: {reward_text}")
    return reward


def read_task_rows() -> list[dict]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, TASK_SHEET_NAME, len(workbook.worksheets) - 1)
    headers = build_header_map(worksheet, TASK_HEADER_ROW)
    required_headers = ["任务id", "任务类型", "奖励内容", "数量", "任务描述", "奖励图标"]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise RuntimeError("Missing task headers: " + ", ".join(missing_headers))

    rows = []
    for row_index in range(TASK_DATA_START_ROW, worksheet.max_row + 1):
        task_id_value = cell_by_header(worksheet, row_index, headers, "任务id")
        task_type_id_value = cell_by_header(worksheet, row_index, headers, "任务类型")
        reward_value = cell_by_header(worksheet, row_index, headers, "奖励内容")
        amount_value = cell_by_header(worksheet, row_index, headers, "数量")
        description_value = cell_by_header(worksheet, row_index, headers, "任务描述")
        icon_value = cell_by_header(worksheet, row_index, headers, "奖励图标")

        if all(is_blank(value) for value in (task_id_value, task_type_id_value, reward_value, amount_value, description_value, icon_value)):
            continue

        task_id = math_safe_int(task_id_value, 0)
        task_type_id = math_safe_int(task_type_id_value, 0)
        if task_id <= 0 or task_type_id <= 0:
            raise RuntimeError(f"Invalid task row {row_index}: task id and type id are required")
        if task_type_id not in TASK_TYPE_BY_ID:
            raise RuntimeError(f"Unknown task type id {task_type_id} at row {row_index}")

        period, task_type = TASK_TYPE_BY_ID[task_type_id]
        description = "" if is_blank(description_value) else str(description_value).strip()
        reward = parse_task_reward(reward_value, amount_value)
        row = {
            "TaskId": task_id,
            "Period": period,
            "TaskType": task_type,
            "TaskTypeId": task_type_id,
            "Target": parse_task_target(task_type, description),
            "RewardType": reward["RewardType"],
            "PotionId": math_safe_int(reward.get("PotionId"), 0),
            "Amount": math_safe_int(reward.get("Amount"), 1),
            "Description": description,
            "Icon": "" if is_blank(icon_value) else str(icon_value).strip(),
        }
        rows.append(row)

    rows.sort(key=lambda row: row["TaskId"])
    return rows


def build_task_generated_block(rows: list[dict]) -> str:
    lines = [
        TASK_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 任务系统数据表. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "TaskConfig.Tasks = {",
    ]
    for row in rows:
        fields = [
            f"TaskId = {row['TaskId']}",
            f"Period = {lua_value(row['Period'])}",
            f"TaskType = {lua_value(row['TaskType'])}",
            f"TaskTypeId = {row['TaskTypeId']}",
            f"Target = {row['Target']}",
            f"RewardType = {lua_value(row['RewardType'])}",
        ]
        if row["PotionId"] > 0:
            fields.append(f"PotionId = {row['PotionId']}")
        fields.extend([
            f"Amount = {row['Amount']}",
            f"Description = {lua_value(row['Description'])}",
            f"Icon = {lua_value(row['Icon'])}",
        ])
        lines.append("    { " + ", ".join(fields) + " },")
    lines.extend(["}", TASK_END_MARKER])
    return "\n".join(lines)


def parse_title_unlock_condition(value) -> tuple[dict | None, str | None]:
    if is_blank(value):
        return None, "Empty title unlock condition"

    text = str(value).strip()
    compact = re.sub(r"\s+", "", text)
    normalized = compact.lower().replace(",", "")
    patterns = [
        (r"(?:玩家)?历史最高等级达到(\d+)级?", "HighestLevelReached"),
        (r"reach(?:lv\.?|level)(\d+)", "HighestLevelReached"),
        (r"(?:玩家)?累计击杀(\d+)人?", "TotalPlayerKills"),
        (r"defeat(\d+)players?intotal\.?", "TotalPlayerKills"),
        (r"(?:玩家)?累计死亡(\d+)次?", "TotalDeaths"),
        (r"die(\d+)times?intotal\.?", "TotalDeaths"),
        (r"(?:玩家)?累计获得(?:钻石|宝石)(\d+)", "TotalDiamondsEarned"),
        (r"earn(\d+)(?:gems?|diamonds?)intotal\.?", "TotalDiamondsEarned"),
        (r"(?:玩家)?累计在线时长达到(\d+)小时", "TotalOnlineHours"),
        (r"(?:stayonline|beonline|online|play)(?:for)?(\d+)hours?(?:intotal)?\.?", "TotalOnlineHours"),
    ]
    for pattern, condition_type in patterns:
        match = re.fullmatch(pattern, normalized)
        if match:
            return {
                "Type": condition_type,
                "Target": math_safe_int(match.group(1), 0),
            }, None

    return None, f"Unable to parse title unlock condition: {text}"


def build_reward_label(reward: dict, workbook_metadata: dict[int, dict]) -> str:
    if not is_blank(reward.get("Label")):
        return str(reward["Label"])
    reward_type = reward.get("RewardType")
    amount = math_safe_int(reward.get("Amount"), 1)
    if reward_type == "WheelSpins":
        return f"Spin x{amount}"
    if reward_type == "Potion":
        potion_id = math_safe_int(reward.get("PotionId"), 0)
        potion_names = {
            1001: "Basic Potion",
            1002: "Advanced Potion",
            1003: "Rare Potion",
        }
        return f"{potion_names.get(potion_id, 'Potion')} x{amount}"
    if reward_type == "Skin":
        skin_id = math_safe_int(reward.get("SkinId"), 0)
        return workbook_metadata.get(skin_id, {}).get("Name") or f"Skin {skin_id}"
    if reward_type == "Diamonds":
        return f"Diamonds x{amount}"
    if reward_type == "Experience":
        return f"EXP x{amount}"
    return f"{reward_type} x{amount}"


def build_reward_icon(reward: dict, workbook_metadata: dict[int, dict]) -> str:
    if not is_blank(reward.get("Icon")):
        return str(reward["Icon"]).strip()
    reward_type = reward.get("RewardType")
    if reward_type == "WheelSpins":
        return "rbxassetid://77152368516350"
    if reward_type == "Potion":
        potion_icons = {
            1001: "rbxassetid://111415582573034",
            1002: "rbxassetid://106498508152369",
            1003: "rbxassetid://100154459165982",
        }
        return potion_icons.get(math_safe_int(reward.get("PotionId"), 0), "")
    if reward_type == "Skin":
        return workbook_metadata.get(math_safe_int(reward.get("SkinId"), 0), {}).get("IconImage") or ""
    if reward_type == "Diamonds":
        return "rbxassetid://89590364394067"
    if reward_type == "Experience":
        return "rbxassetid://112367399278116"
    return ""


def build_seven_day_reward_row(reward: dict, workbook_metadata: dict[int, dict]) -> str:
    normalized = dict(reward)
    normalized["Label"] = build_reward_label(normalized, workbook_metadata)
    normalized["Icon"] = build_reward_icon(normalized, workbook_metadata)
    fields = [
        f"DayIndex = {math_safe_int(normalized.get('DayIndex'), 1)}",
        f"RewardType = {lua_value(normalized.get('RewardType'))}",
    ]
    if normalized.get("PotionId"):
        fields.append(f"PotionId = {math_safe_int(normalized.get('PotionId'), 0)}")
    if normalized.get("SkinId"):
        fields.append(f"SkinId = {math_safe_int(normalized.get('SkinId'), 0)}")
    fields.append(f"Amount = {math_safe_int(normalized.get('Amount'), 1)}")
    fields.append(f"Label = {lua_value(normalized.get('Label'))}")
    fields.append(f"Icon = {lua_value(normalized.get('Icon'))}")
    return "    { " + ", ".join(fields) + " },"


def read_seven_day_login_reward_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, SEVEN_DAY_LOGIN_REWARD_SHEET_NAME, 16)
    metadata = read_weapon_skin_metadata(workbook)
    warnings = []

    def read_block(start_row: int, end_row: int) -> list[dict]:
        rows = []
        header_values = [str(worksheet.cell(start_row - 1, column).value or "").strip() for column in range(1, worksheet.max_column + 1)]
        label_column = None
        icon_column = None
        for index, header in enumerate(header_values, start=1):
            if header == "奖励名字":
                label_column = index
            elif header == "奖励图标":
                icon_column = index
        for row_index in range(start_row, end_row + 1):
            day_value = worksheet.cell(row_index, 4).value
            reward_type_value = worksheet.cell(row_index, 5).value
            amount_value = worksheet.cell(row_index, 6).value
            if all(is_blank(value) for value in (day_value, reward_type_value, amount_value)):
                continue
            reward, warning = parse_reward_type_and_amount(reward_type_value, amount_value)
            if warning:
                warnings.append({"row": row_index, "column": 5, "warning": warning})
            if not reward:
                continue
            reward["DayIndex"] = math_safe_int(day_value, len(rows) + 1)
            if label_column:
                reward["Label"] = "" if is_blank(worksheet.cell(row_index, label_column).value) else str(worksheet.cell(row_index, label_column).value).strip()
            if icon_column:
                reward["Icon"] = "" if is_blank(worksheet.cell(row_index, icon_column).value) else str(worksheet.cell(row_index, icon_column).value).strip()
            rows.append(reward)
        return rows

    first_cycle_rows = read_block(6, 12)
    repeat_cycle_rows = read_block(20, 26)
    return first_cycle_rows, repeat_cycle_rows, metadata, warnings


def build_seven_day_login_reward_generated_block(first_cycle_rows, repeat_cycle_rows, metadata) -> str:
    lines = [
        SEVEN_DAY_LOGIN_REWARD_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 七日登录奖励. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "SevenDayLoginRewardConfig.FirstCycleRewards = {",
    ]
    for reward in first_cycle_rows:
        lines.append(build_seven_day_reward_row(reward, metadata))
    lines.extend(["}", "", "SevenDayLoginRewardConfig.RepeatCycleRewards = {"])
    for reward in repeat_cycle_rows:
        lines.append(build_seven_day_reward_row(reward, metadata))
    lines.extend(["}", SEVEN_DAY_LOGIN_REWARD_END_MARKER])
    return "\n".join(lines)


def read_skin_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    skin_sheet = get_sheet(workbook, "皮肤表", 12)
    metadata = read_weapon_skin_metadata(workbook)
    rows = []
    for row_index in range(5, skin_sheet.max_row + 1):
        skin_id = math_safe_int(skin_sheet.cell(row_index, 3).value, 0)
        if skin_id <= 0:
            continue
        channel = math_safe_int(skin_sheet.cell(row_index, 4).value, 0)
        if channel <= 0:
            continue
        info = metadata.get(skin_id)
        if not info:
            continue
        rows.append({
            "Id": skin_id,
            "Name": info["Name"],
            "TemplateName": info["TemplateName"],
            "TemplatePath": info["TemplatePath"],
            "IconImage": info["IconImage"],
            "PurchaseChannel": channel,
            "DiamondPrice": math_safe_int(skin_sheet.cell(row_index, 5).value, 0),
            "RobuxPrice": math_safe_int(skin_sheet.cell(row_index, 6).value, 0),
            "GamePassId": 1830742687 if skin_id == 10002 else 0,
        })
    return rows


def build_skin_generated_block(rows) -> str:
    channel_names = {
        1: "Diamonds",
        2: "GamePass",
        3: "Wheel",
        4: "SevenDayLoginReward",
    }
    lines = [
        SKIN_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 皮肤表 + 武器数值. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "SkinConfig.Skins = {",
    ]
    for row in rows:
        channel_name = channel_names.get(row["PurchaseChannel"], "Diamonds")
        lines.extend([
            "    {",
            f"        Id = {row['Id']},",
            f"        Name = {lua_value(row['Name'])},",
            f"        TemplateName = {lua_value(row['TemplateName'])},",
            f"        TemplatePath = {lua_value(row['TemplatePath'])},",
            f"        IconImage = {lua_value(row['IconImage'])},",
            f"        PurchaseChannel = SkinConfig.PurchaseChannel.{channel_name},",
            f"        DiamondPrice = {row['DiamondPrice']},",
            f"        GamePassId = {row['GamePassId']},",
            "    },",
        ])
    lines.extend(["}", SKIN_END_MARKER])
    return "\n".join(lines)


def read_trail_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, TRAIL_SHEET_NAME, len(workbook.worksheets) - 1)
    rows = []
    for row_index in range(TRAIL_DATA_START_ROW, worksheet.max_row + 1):
        trail_id = math_safe_int(worksheet.cell(row_index, 3).value, 0)
        if trail_id <= 0:
            continue
        trail_name = worksheet.cell(row_index, 4).value
        template_name = worksheet.cell(row_index, 5).value
        if is_blank(template_name):
            continue
        rows.append({
            "Id": trail_id,
            "Name": "" if is_blank(trail_name) else str(trail_name).strip(),
            "TemplateName": str(template_name).strip(),
            "TemplatePath": "ReplicatedStorage/Model/Trail/" + str(template_name).strip(),
            "DiamondPrice": math_safe_int(worksheet.cell(row_index, 6).value, 0),
            "RobuxPrice": math_safe_int(worksheet.cell(row_index, 7).value, 0),
            "ProductId": math_safe_int(worksheet.cell(row_index, 8).value, 0),
            "IsDefaultUnlocked": math_safe_int(worksheet.cell(row_index, 9).value, 0) == 1,
            "IconImage": "" if is_blank(worksheet.cell(row_index, 10).value) else str(worksheet.cell(row_index, 10).value).strip(),
        })
    return rows


def build_trail_generated_block(rows) -> str:
    lines = [
        TRAIL_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 尾迹. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "TrailConfig.Trails = {",
    ]
    for row in rows:
        lines.extend([
            "    {",
            f"        Id = {row['Id']},",
            f"        Name = {lua_value(row['Name'])},",
            f"        TemplateName = {lua_value(row['TemplateName'])},",
            f"        TemplatePath = {lua_value(row['TemplatePath'])},",
            f"        IconImage = {lua_value(row['IconImage'])},",
            f"        DiamondPrice = {row['DiamondPrice']},",
            f"        RobuxPrice = {row['RobuxPrice']},",
            f"        ProductId = {row['ProductId']},",
            f"        IsDefaultUnlocked = {'true' if row['IsDefaultUnlocked'] else 'false'},",
            "    },",
        ])
    lines.extend(["}", TRAIL_END_MARKER])
    return "\n".join(lines)


def read_title_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, TITLE_SHEET_NAME, len(workbook.worksheets) - 1)
    header_by_name = {}
    for column_index in range(1, worksheet.max_column + 1):
        header = worksheet.cell(TITLE_HEADER_ROW, column_index).value
        if not is_blank(header):
            header_by_name[str(header).strip()] = column_index

    required_headers = ["称号Id", "称号名字", "称号描述", "称号解锁条件", "称号图片资源"]
    missing_headers = [header for header in required_headers if header not in header_by_name]
    if missing_headers:
        raise RuntimeError("Missing title sheet headers: " + ", ".join(missing_headers))

    rows = []
    warnings = []
    for row_index in range(TITLE_DATA_START_ROW, worksheet.max_row + 1):
        title_id = math_safe_int(worksheet.cell(row_index, header_by_name["称号Id"]).value, 0)
        if title_id <= 0:
            continue

        name = worksheet.cell(row_index, header_by_name["称号名字"]).value
        description = worksheet.cell(row_index, header_by_name["称号描述"]).value
        condition_text = worksheet.cell(row_index, header_by_name["称号解锁条件"]).value
        icon_image = worksheet.cell(row_index, header_by_name["称号图片资源"]).value
        condition, warning = parse_title_unlock_condition(condition_text)
        if warning:
            warnings.append(f"row {row_index}: {warning}")

        rows.append({
            "Id": title_id,
            "Name": "" if is_blank(name) else str(name).strip(),
            "Description": "" if is_blank(description) else str(description).strip(),
            "UnlockConditionText": "" if is_blank(condition_text) else str(condition_text).strip(),
            "IconImage": "" if is_blank(icon_image) else str(icon_image).strip(),
            "Condition": condition,
        })
    return rows, warnings


def build_title_generated_block(rows) -> str:
    lines = [
        TITLE_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 称号. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "TitleConfig.Titles = {",
    ]
    for row in rows:
        lines.extend([
            "    {",
            f"        Id = {row['Id']},",
            f"        Name = {lua_value(row['Name'])},",
            f"        Description = {lua_value(row['Description'])},",
            f"        UnlockConditionText = {lua_value(row['UnlockConditionText'])},",
            f"        IconImage = {lua_value(row['IconImage'])},",
        ])
        condition = row.get("Condition")
        if condition:
            lines.append(
                "        Condition = { Type = %s, Target = %d },"
                % (lua_value(condition["Type"]), math_safe_int(condition["Target"], 0))
            )
        lines.extend([
            "    },",
        ])
    lines.extend(["}", TITLE_END_MARKER])
    return "\n".join(lines)


def read_attribute_config_rows():
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    config_sheet = get_sheet(workbook, ATTRIBUTE_CONFIG_SHEET_NAME, len(workbook.worksheets) - 3)
    product_sheet = get_sheet(workbook, ATTRIBUTE_PRODUCT_SHEET_NAME, len(workbook.worksheets) - 1)
    config_headers = build_header_map(config_sheet, ATTRIBUTE_CONFIG_HEADER_ROW)
    product_headers = build_header_map(product_sheet, ATTRIBUTE_PRODUCT_HEADER_ROW)

    required_config_headers = [
        "AttributeID",
        "英文显示名",
        "上限显示名",
        "CardName",
        "初始上限",
        "最高上限",
        "单局每级效果值",
        "显示单位/ValueType",
        "DeveloperProductId",
        "钻石购买",
        "罗布币购买",
    ]
    missing_config_headers = [header for header in required_config_headers if header not in config_headers]
    if missing_config_headers:
        raise RuntimeError("Missing attribute config headers: " + ", ".join(missing_config_headers))

    required_product_headers = [
        "等级",
        "开发者商品",
        "钻石价格",
    ]
    missing_product_headers = [header for header in required_product_headers if header not in product_headers]
    if missing_product_headers:
        raise RuntimeError("Missing attribute product headers: " + ", ".join(missing_product_headers))

    attributes = []
    for row_index in range(ATTRIBUTE_CONFIG_DATA_START_ROW, config_sheet.max_row + 1):
        attribute_id = cell_by_header(config_sheet, row_index, config_headers, "AttributeID")
        if is_blank(attribute_id):
            continue

        attribute_key = str(attribute_id).strip()
        attributes.append({
            "Key": attribute_key,
            "DisplayName": "" if is_blank(cell_by_header(config_sheet, row_index, config_headers, "英文显示名")) else str(cell_by_header(config_sheet, row_index, config_headers, "英文显示名")).strip(),
            "CapDisplayName": "" if is_blank(cell_by_header(config_sheet, row_index, config_headers, "上限显示名")) else str(cell_by_header(config_sheet, row_index, config_headers, "上限显示名")).strip(),
            "CardName": "" if is_blank(cell_by_header(config_sheet, row_index, config_headers, "CardName")) else str(cell_by_header(config_sheet, row_index, config_headers, "CardName")).strip(),
            "InitialCap": math_safe_int(cell_by_header(config_sheet, row_index, config_headers, "初始上限"), 0),
            "MaxCap": math_safe_int(cell_by_header(config_sheet, row_index, config_headers, "最高上限"), 0),
            "PerLevelValue": math_safe_float(cell_by_header(config_sheet, row_index, config_headers, "单局每级效果值"), 0),
            "ValueType": normalize_attribute_value_type(cell_by_header(config_sheet, row_index, config_headers, "显示单位/ValueType")),
            "ProductId": math_safe_int(cell_by_header(config_sheet, row_index, config_headers, "DeveloperProductId"), 0),
            "GemEnabled": bool_from_cell(cell_by_header(config_sheet, row_index, config_headers, "钻石购买"), True),
            "RobuxEnabled": bool_from_cell(cell_by_header(config_sheet, row_index, config_headers, "罗布币购买"), True),
        })

    level_products = []
    common_price_rows = []
    for row_index in range(ATTRIBUTE_PRODUCT_DATA_START_ROW, product_sheet.max_row + 1):
        level = cell_by_header(product_sheet, row_index, product_headers, "等级")
        product_id = cell_by_header(product_sheet, row_index, product_headers, "开发者商品")
        gem_cost = cell_by_header(product_sheet, row_index, product_headers, "钻石价格")
        if not isinstance(level, (int, float)) or not isinstance(product_id, (int, float)) or not isinstance(gem_cost, (int, float)):
            continue
        resolved_level = math_safe_int(level, 0)
        resolved_product_id = math_safe_int(product_id, 0)
        resolved_gem_cost = math_safe_int(gem_cost, 0)
        if resolved_level <= 0 or resolved_product_id <= 0:
            continue
        level_products.append({
            "Level": resolved_level,
            "ProductId": resolved_product_id,
            "GemCost": resolved_gem_cost,
        })
        if resolved_gem_cost > 0:
            common_price_rows.append({
                "FromCap": resolved_level - 1,
                "ToCap": resolved_level,
                "GemCost": resolved_gem_cost,
            })
    level_products.sort(key=lambda row: row["Level"])
    common_price_rows.sort(key=lambda row: (row["FromCap"], row["ToCap"]))

    prices_by_key = {}
    for attribute in attributes:
        initial_cap = math_safe_int(attribute["InitialCap"], 0)
        max_cap = math_safe_int(attribute["MaxCap"], initial_cap)
        prices_by_key[attribute["Key"]] = []
        for row in common_price_rows:
            if row["FromCap"] < initial_cap or row["ToCap"] > max_cap:
                continue
            prices_by_key[attribute["Key"]].append({
                **row,
                "GemEnabled": attribute["GemEnabled"],
                "RobuxEnabled": attribute["RobuxEnabled"],
            })

    return attributes, prices_by_key, level_products


def normalize_attribute_value_type(value) -> str:
    text = "" if is_blank(value) else str(value).strip()
    if text == "%":
        return "Percent"
    return text or "Percent"


def lua_number(value) -> str:
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return ("%0.12g" % number)


def build_attribute_config_generated_block(attributes: list[dict], prices_by_key: dict[str, list[dict]], level_products: list[dict]) -> str:
    lines = [
        ATTRIBUTE_CONFIG_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 属性养成配置 + 属性养成新的开发者商品. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "AttributeConfig.Order = {",
    ]
    for row in attributes:
        lines.append(f"    {lua_value(row['Key'])},")
    lines.extend(["}", "", "AttributeConfig.Attributes = {"])
    for row in attributes:
        lines.extend([
            f"    {row['Key']} = {{",
            f"        DisplayName = {lua_value(row['DisplayName'])},",
            f"        CapDisplayName = {lua_value(row['CapDisplayName'])},",
            f"        CardName = {lua_value(row['CardName'])},",
            f"        InitialCap = {row['InitialCap']},",
            f"        MaxCap = {row['MaxCap']},",
            f"        PerLevelValue = {lua_number(row['PerLevelValue'])},",
            f"        ValueType = {lua_value(row['ValueType'])},",
            "    },",
        ])
    lines.extend(["}", "", "AttributeConfig.CapUpgradeProducts = {}", "", "AttributeConfig.CapUpgradeLevelProducts = {"])
    for row in level_products:
        lines.append(f"    [{row['Level']}] = {row['ProductId']},")
    lines.extend(["}", "", "AttributeConfig.CapUpgradePrices = {"])
    for row in attributes:
        key = row["Key"]
        lines.append(f"    {key} = {{")
        for price in prices_by_key.get(key, []):
            lines.append(
                "        { FromCap = %d, ToCap = %d, GemCost = %d, GemEnabled = %s, RobuxEnabled = %s },"
                % (
                    price["FromCap"],
                    price["ToCap"],
                    price["GemCost"],
                    "true" if price["GemEnabled"] else "false",
                    "true" if price["RobuxEnabled"] else "false",
                )
            )
        lines.append("    },")
    lines.extend(["}", ATTRIBUTE_CONFIG_END_MARKER])
    return "\n".join(lines)


def read_diamond_shop_rows() -> list[dict]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    worksheet = get_sheet(workbook, DIAMOND_SHOP_SHEET_NAME, len(workbook.worksheets) - 1)
    headers = build_header_map(worksheet, DIAMOND_SHOP_HEADER_ROW)

    required_headers = ["id", "开发者商品id", "钻石数"]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise RuntimeError("Missing diamond shop headers: " + ", ".join(missing_headers))

    rows = []
    for row_index in range(DIAMOND_SHOP_DATA_START_ROW, worksheet.max_row + 1):
        diamond_id = cell_by_header(worksheet, row_index, headers, "id")
        product_id = cell_by_header(worksheet, row_index, headers, "开发者商品id")
        diamonds = cell_by_header(worksheet, row_index, headers, "钻石数")
        if all(is_blank(value) for value in (diamond_id, product_id, diamonds)):
            continue

        resolved_id = math_safe_int(diamond_id, 0)
        resolved_product_id = math_safe_int(product_id, 0)
        resolved_diamonds = math_safe_int(diamonds, 0)
        if resolved_id <= 0 or resolved_product_id <= 0 or resolved_diamonds <= 0:
            raise RuntimeError(f"Invalid diamond shop row {row_index}")

        rows.append({
            "Id": resolved_id,
            "ProductId": resolved_product_id,
            "Diamonds": resolved_diamonds,
        })

    rows.sort(key=lambda row: row["Id"])
    return rows


def build_diamond_shop_generated_block(rows: list[dict]) -> str:
    lines = [
        DIAMOND_SHOP_BEGIN_MARKER,
        "-- Source: IO_BaseBalanceDraft.xlsx / 钻石购买. Update via tools/SyncCodeConfigFromWorkbook.py.",
        "ShopConfig.DiamondProducts = {",
    ]
    for row in rows:
        lines.append(
            "    { Id = %d, ProductId = %d, Diamonds = %d },"
            % (row["Id"], row["ProductId"], row["Diamonds"])
        )
    lines.extend(["}", DIAMOND_SHOP_END_MARKER])
    return "\n".join(lines)


def replace_generated_block(source: str, generated_block: str, begin_marker: str, end_marker: str, config_path: Path) -> str:
    pattern = re.compile(re.escape(begin_marker) + r".*?" + re.escape(end_marker), re.S)
    if not pattern.search(source):
        raise RuntimeError(f"Could not find generated block markers in {config_path}")
    return pattern.sub(generated_block, source, count=1)


def sync_attribute_config() -> dict:
    attribute_rows, attribute_price_rows, attribute_level_product_rows = read_attribute_config_rows()
    attribute_source = ATTRIBUTE_CONFIG_PATH.read_text(encoding="utf-8")
    updated_attribute_source = replace_generated_block(
        attribute_source,
        build_attribute_config_generated_block(attribute_rows, attribute_price_rows, attribute_level_product_rows),
        ATTRIBUTE_CONFIG_BEGIN_MARKER,
        ATTRIBUTE_CONFIG_END_MARKER,
        ATTRIBUTE_CONFIG_PATH,
    )
    ATTRIBUTE_CONFIG_PATH.write_text(updated_attribute_source, encoding="utf-8", newline="\n")
    return {
        "attributeRows": len(attribute_rows),
        "attributePriceRows": sum(len(rows) for rows in attribute_price_rows.values()),
        "attributeProductRows": len(attribute_level_product_rows),
    }


def sync_diamond_shop_config() -> dict:
    diamond_rows = read_diamond_shop_rows()
    shop_source = SHOP_CONFIG_PATH.read_text(encoding="utf-8")
    updated_shop_source = replace_generated_block(
        shop_source,
        build_diamond_shop_generated_block(diamond_rows),
        DIAMOND_SHOP_BEGIN_MARKER,
        DIAMOND_SHOP_END_MARKER,
        SHOP_CONFIG_PATH,
    )
    SHOP_CONFIG_PATH.write_text(updated_shop_source, encoding="utf-8", newline="\n")
    return {
        "diamondShopRows": len(diamond_rows),
    }


def sync_potion_config() -> dict:
    potion_rows = read_potion_rows()
    potion_source = POTION_CONFIG_PATH.read_text(encoding="utf-8")
    updated_potion_source = replace_generated_block(
        potion_source,
        build_potion_generated_block(potion_rows),
        POTION_BEGIN_MARKER,
        POTION_END_MARKER,
        POTION_CONFIG_PATH,
    )
    POTION_CONFIG_PATH.write_text(updated_potion_source, encoding="utf-8", newline="\n")
    return {
        "potionRows": len(potion_rows),
    }


def sync_wheel_config() -> dict:
    wheel_rows, wheel_warnings = read_wheel_reward_rows()
    wheel_source = WHEEL_CONFIG_PATH.read_text(encoding="utf-8")
    updated_wheel_source = replace_generated_block(
        wheel_source,
        build_wheel_reward_generated_block(wheel_rows),
        WHEEL_BEGIN_MARKER,
        WHEEL_END_MARKER,
        WHEEL_CONFIG_PATH,
    )
    WHEEL_CONFIG_PATH.write_text(updated_wheel_source, encoding="utf-8", newline="\n")
    return {
        "wheelRewardRows": len(wheel_rows),
        "warnings": wheel_warnings,
    }


def sync_seven_day_login_reward_config() -> dict:
    first_cycle_rows, repeat_cycle_rows, skin_metadata, seven_day_warnings = read_seven_day_login_reward_rows()
    seven_day_source = SEVEN_DAY_LOGIN_REWARD_CONFIG_PATH.read_text(encoding="utf-8")
    updated_seven_day_source = replace_generated_block(
        seven_day_source,
        build_seven_day_login_reward_generated_block(first_cycle_rows, repeat_cycle_rows, skin_metadata),
        SEVEN_DAY_LOGIN_REWARD_BEGIN_MARKER,
        SEVEN_DAY_LOGIN_REWARD_END_MARKER,
        SEVEN_DAY_LOGIN_REWARD_CONFIG_PATH,
    )
    SEVEN_DAY_LOGIN_REWARD_CONFIG_PATH.write_text(updated_seven_day_source, encoding="utf-8", newline="\n")
    return {
        "sevenDayFirstCycleRows": len(first_cycle_rows),
        "sevenDayRepeatCycleRows": len(repeat_cycle_rows),
        "warnings": seven_day_warnings,
    }


def sync_title_config() -> dict:
    title_rows, title_warnings = read_title_rows()
    title_source = TITLE_CONFIG_PATH.read_text(encoding="utf-8")
    updated_title_source = replace_generated_block(
        title_source,
        build_title_generated_block(title_rows),
        TITLE_BEGIN_MARKER,
        TITLE_END_MARKER,
        TITLE_CONFIG_PATH,
    )
    TITLE_CONFIG_PATH.write_text(updated_title_source, encoding="utf-8", newline="\n")
    return {
        "titleRows": len(title_rows),
        "warnings": title_warnings,
    }


def sync_task_config() -> dict:
    task_rows = read_task_rows()
    task_source = TASK_CONFIG_PATH.read_text(encoding="utf-8")
    updated_task_source = replace_generated_block(
        task_source,
        build_task_generated_block(task_rows),
        TASK_BEGIN_MARKER,
        TASK_END_MARKER,
        TASK_CONFIG_PATH,
    )
    TASK_CONFIG_PATH.write_text(updated_task_source, encoding="utf-8", newline="\n")
    return {
        "taskRows": len(task_rows),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Sync Lua config files from IO_BaseBalanceDraft.xlsx.")
    parser.add_argument("--attribute-only", action="store_true", help="Only sync AttributeConfig.lua from attribute progression sheets.")
    parser.add_argument("--diamond-shop-only", action="store_true", help="Only sync ShopConfig.lua diamond products from the diamond purchase sheet.")
    parser.add_argument("--seven-day-only", action="store_true", help="Only sync SevenDayLoginRewardConfig.lua from the seven-day login reward sheet.")
    parser.add_argument("--title-only", action="store_true", help="Only sync TitleConfig.lua from the title sheet.")
    parser.add_argument("--task-only", action="store_true", help="Only sync TaskConfig.lua from the task sheet.")
    args = parser.parse_args()

    if args.diamond_shop_only:
        print(json.dumps({
            **sync_diamond_shop_config(),
            "warnings": [],
        }, ensure_ascii=False))
        return
    if args.seven_day_only:
        print(json.dumps(sync_seven_day_login_reward_config(), ensure_ascii=False))
        return
    if args.title_only:
        print(json.dumps(sync_title_config(), ensure_ascii=False))
        return
    if args.task_only:
        print(json.dumps({
            **sync_task_config(),
            "warnings": [],
        }, ensure_ascii=False))
        return

    attribute_result = sync_attribute_config()
    if args.attribute_only:
        print(json.dumps({
            **attribute_result,
            "warnings": [],
        }, ensure_ascii=False))
        return

    potion_result = sync_potion_config()

    rows, warnings = read_code_rows()
    source = CONFIG_PATH.read_text(encoding="utf-8")
    updated_source = replace_generated_block(source, build_generated_block(rows), BEGIN_MARKER, END_MARKER, CONFIG_PATH)
    CONFIG_PATH.write_text(updated_source, encoding="utf-8", newline="\n")

    online_rows, online_warnings = read_online_reward_rows()
    online_source = ONLINE_REWARD_CONFIG_PATH.read_text(encoding="utf-8")
    updated_online_source = replace_generated_block(
        online_source,
        build_online_reward_generated_block(online_rows),
        ONLINE_REWARD_BEGIN_MARKER,
        ONLINE_REWARD_END_MARKER,
        ONLINE_REWARD_CONFIG_PATH,
    )
    ONLINE_REWARD_CONFIG_PATH.write_text(updated_online_source, encoding="utf-8", newline="\n")

    seven_day_result = sync_seven_day_login_reward_config()

    skin_rows = read_skin_rows()
    skin_source = SKIN_CONFIG_PATH.read_text(encoding="utf-8")
    updated_skin_source = replace_generated_block(
        skin_source,
        build_skin_generated_block(skin_rows),
        SKIN_BEGIN_MARKER,
        SKIN_END_MARKER,
        SKIN_CONFIG_PATH,
    )
    SKIN_CONFIG_PATH.write_text(updated_skin_source, encoding="utf-8", newline="\n")

    trail_rows = read_trail_rows()
    trail_source = TRAIL_CONFIG_PATH.read_text(encoding="utf-8")
    updated_trail_source = replace_generated_block(
        trail_source,
        build_trail_generated_block(trail_rows),
        TRAIL_BEGIN_MARKER,
        TRAIL_END_MARKER,
        TRAIL_CONFIG_PATH,
    )
    TRAIL_CONFIG_PATH.write_text(updated_trail_source, encoding="utf-8", newline="\n")

    title_result = sync_title_config()
    diamond_shop_result = sync_diamond_shop_config()
    wheel_result = sync_wheel_config()
    task_result = sync_task_config()
    print(json.dumps({
        "codeRows": len(rows),
        "onlineRewardRows": len(online_rows),
        "sevenDayFirstCycleRows": seven_day_result["sevenDayFirstCycleRows"],
        "sevenDayRepeatCycleRows": seven_day_result["sevenDayRepeatCycleRows"],
        "skinRows": len(skin_rows),
        "trailRows": len(trail_rows),
        "titleRows": title_result["titleRows"],
        **potion_result,
        **wheel_result,
        **attribute_result,
        **diamond_shop_result,
        **task_result,
        "warnings": warnings + online_warnings + seven_day_result["warnings"] + title_result["warnings"] + wheel_result["warnings"],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
